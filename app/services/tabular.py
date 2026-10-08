"""Чтение и запись таблиц CSV / XLSX без привязки к базе данных.

Чтение: файл -> Table(заголовки, строки). Запись: заголовки + строки -> байты файла.
CSV пишется в UTF-8 с BOM и разделителем «;» — так его без плясок открывает русский Excel;
при чтении кодировка (UTF-8 / Windows-1251) и разделитель определяются автоматически."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv; charset=utf-8"
FORMATS = ("xlsx", "csv")
MAX_VALIDATION_ROWS = 2000      # на сколько строк шаблона действуют выпадающие списки
CSV_DELIMITERS = ";,\t"


class FileError(Exception):
    """Файл нельзя прочитать вообще; текст показывается пользователю."""


@dataclass
class Table:
    headers: list[str]
    rows: list[tuple[int, list[Any]]]  # (номер строки в файле, значения ячеек)


# ---------------------------------------------------------------------- чтение
def normalize_header(value: Any) -> str:
    """Заголовок для сопоставления: без регистра, «ё» = «е», без «*» и лишних пробелов."""
    text = str(value or "").replace("*", " ").replace("ё", "е").replace("Ё", "Е")
    return " ".join(text.split()).casefold()


def _clean_cell(value: Any) -> Any:
    """Приводит значение ячейки к простому виду: пустое -> None, 5.0 -> 5, строка без пробелов по краям."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        # так экспорт защищает текст, похожий на формулу (см. safe_text)
        if len(text) > 1 and text[0] == "'" and text[1] in "=+-@":
            text = text[1:]
        return text or None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return value


def read_table(data: bytes, filename: str, *, max_rows: int) -> Table:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        raw_rows = _read_xlsx(data)
    elif name.endswith(".csv") or name.endswith(".txt"):
        raw_rows = _read_csv(data)
    elif name.endswith(".xls"):
        raise FileError("Старый формат .xls не поддерживается: сохраните файл как .xlsx или .csv")
    else:
        raise FileError("Поддерживаются файлы .xlsx и .csv")

    headers: list[str] | None = None
    rows: list[tuple[int, list[Any]]] = []
    for number, cells in raw_rows:
        cells = [_clean_cell(c) for c in cells]
        if all(c is None for c in cells):
            continue
        if headers is None:
            headers = ["" if c is None else str(c).strip() for c in cells]
            continue
        rows.append((number, cells))
        if len(rows) > max_rows:
            raise FileError(f"Слишком много строк: в одном файле можно загрузить до {max_rows}")
    if headers is None:
        raise FileError("Файл пустой: в нём нет строки с заголовками")
    # лишние пустые колонки справа не нужны
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise FileError("В первой строке файла нет заголовков колонок")
    return Table(headers=headers, rows=rows)


def _read_xlsx(data: bytes):
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — openpyxl бросает разные исключения на битых файлах
        raise FileError("Не удалось прочитать файл Excel. Проверьте, что это файл .xlsx") from None
    try:
        sheet = wb.worksheets[0] if wb.worksheets else None
        if sheet is None:
            raise FileError("В книге Excel нет листов")
        return [(i, list(row)) for i, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    except FileError:
        raise
    except Exception:  # noqa: BLE001
        raise FileError("Не удалось прочитать лист Excel") from None
    finally:
        wb.close()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileError("Не удалось определить кодировку файла. Сохраните CSV в UTF-8")


def _detect_delimiter(text: str) -> str:
    first = next((line for line in text.splitlines() if line.strip()), "")
    counts = {d: first.count(d) for d in CSV_DELIMITERS}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] else ";"


def _read_csv(data: bytes):
    if b"\x00" in data[:4096]:
        raise FileError("Это не текстовый CSV-файл")
    text = _decode(data)
    try:
        reader = csv.reader(io.StringIO(text), delimiter=_detect_delimiter(text))
        return [(i, row) for i, row in enumerate(reader, start=1)]
    except csv.Error as e:
        raise FileError(f"Не удалось разобрать CSV: {e}") from None


# ---------------------------------------------------------------------- запись
def safe_text(value: str) -> str:
    """Защита от «формул» в Excel: текст, начинающийся с = @ или похожий на формулу с + / -, получает апостроф.
    Телефоны вида «+7 (916) 123-45-67» и отрицательные числа остаются как есть."""
    if not value:
        return value
    first = value[0]
    if first in "=@\t\r":
        return "'" + value
    if first in "+-" and not re.fullmatch(r"[+\-]?[\d\s().\-]+", value):
        return "'" + value
    return value


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, Decimal):
        text = format(value.normalize(), "f") if value == value.to_integral_value() else format(value, "f")
        return text.replace(".", ",")
    if isinstance(value, float):
        return repr(value).replace(".", ",")
    return safe_text(str(value))


def write_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([safe_text(h) for h in headers])
    for row in rows:
        writer.writerow([_csv_value(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")


@dataclass
class ListValidation:
    """Выпадающий список в колонке шаблона."""
    column: int               # номер колонки, с 1
    values: list[str]


@dataclass
class XlsxSpec:
    headers: list[str]
    rows: list[list[Any]] = field(default_factory=list)
    sheet_title: str = "Данные"
    hints: list[tuple[str, str, str]] = field(default_factory=list)  # (колонка, обязательна?, что вводить)
    lists: list[ListValidation] = field(default_factory=list)
    money_columns: set[int] = field(default_factory=set)             # номера колонок (с 1) с суммами


_HEADER_FILL = PatternFill("solid", fgColor="E0E7FF")


def _cell_value(value: Any) -> Any:
    """Значение для ячейки Excel; строки-«формулы» получают апостроф, остальное пишется как есть."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        return safe_text(value)
    return value


def write_xlsx(spec: XlsxSpec) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_title
    ws.append([safe_text(h) for h in spec.headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    widths = [min(max(len(h) + 2, 12), 40) for h in spec.headers]
    for row in spec.rows:
        values = [_cell_value(v) for v in row]
        ws.append(values)
        for i, v in enumerate(values):
            if isinstance(v, str):
                widths[i] = min(max(widths[i], len(v) + 2), 50)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
            elif isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
            elif cell.column in spec.money_columns and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, str):
                # строка, начинающаяся с «=», не должна стать формулой
                cell.data_type = "s"
    if spec.rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(spec.headers))}{len(spec.rows) + 1}"

    if spec.lists:
        lists_ws = wb.create_sheet("Списки")
        for position, item in enumerate(spec.lists, start=1):
            lists_ws.cell(row=1, column=position, value=safe_text(spec.headers[item.column - 1]))
            for r, value in enumerate(item.values, start=2):
                lists_ws.cell(row=r, column=position, value=safe_text(value)).data_type = "s"
            if not item.values:
                continue
            letter = get_column_letter(position)
            last = len(item.values) + 1
            validation = DataValidation(
                type="list", formula1=f"'Списки'!${letter}$2:${letter}${last}", allow_blank=True,
                showErrorMessage=True, errorTitle="Недопустимое значение",
                error="Выберите значение из списка.",
            )
            ws.add_data_validation(validation)
            col_letter = get_column_letter(item.column)
            validation.add(f"{col_letter}2:{col_letter}{MAX_VALIDATION_ROWS}")
        lists_ws.sheet_state = "hidden"

    if spec.hints:
        hints_ws = wb.create_sheet("Подсказки")
        hints_ws.append(["Колонка", "Обязательна", "Что вводить"])
        for cell in hints_ws[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
        for column, required, text in spec.hints:
            hints_ws.append([safe_text(column), required, text])
        hints_ws.column_dimensions["A"].width = 30
        hints_ws.column_dimensions["B"].width = 14
        hints_ws.column_dimensions["C"].width = 90
        for row in hints_ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
