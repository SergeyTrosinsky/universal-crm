"""Импорт и экспорт клиентов и сделок: CSV / Excel, предпросмотр, ошибки по строкам, права."""
import io
import re
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.config import get_settings
from app.core.timezone import to_local
from app.models import Client, Deal
from app.services import custom_field_service, presets, role_service
from tests.conftest import PASSWORD


def web_login(client, email):
    client.cookies.clear()
    resp = client.post("/login", data={"email": email, "password": PASSWORD})
    assert resp.status_code in (302, 303), resp.text[:300]


def csv_file(rows, delimiter=";"):
    return "\r\n".join(delimiter.join(row) for row in rows).encode("utf-8-sig")


def xlsx_file(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, kind, name, data, **form):
    return client.post(f"/import/{kind}", files={"file": (name, data)}, data=form)


def token_of(resp) -> str:
    match = re.search(r'name="token" value="([0-9a-f]{32})"', resp.text)
    assert match, resp.text[:800]
    return match.group(1)


def confirm(client, kind, token):
    return client.post(f"/import/{kind}/confirm", data={"token": token})


def login_admin(client, make_user, email="admin@example.com"):
    make_user(email, role="admin", full_name="Админ")
    web_login(client, email)


# ------------------------------------------------------------------ клиенты: проверка и загрузка
def test_clients_import_checks_first_then_loads(client, make_user, factory):
    login_admin(client, make_user)
    with factory() as s:
        source = custom_field_service.create_field(
            s, entity_type="client", label="Источник", field_type="select", options=["Сайт", "Звонок"]
        ).code
    data = csv_file(
        [
            ["Имя", "Телефон", "Email", "Тип", "Источник"],
            ["Анна", "+7 (916) 111-22-33", "anna@example.com", "Физлицо", "Сайт"],
            ["ООО Ромашка", "8 495 000-11-22", "", "Компания", ""],
            ["Плохой телефон", "12", "", "", ""],
            ["Плохой тип", "", "", "Робот", ""],
            ["Анна ещё раз", "9161112233", "", "", ""],
            ["Плохой email", "", "x@@bad", "", ""],
            ["Плохой список", "", "", "", "Реклама"],
        ]
    )
    preview = upload(client, "clients", "clients.csv", data)
    assert preview.status_code == 200
    assert "Ошибки (4)" in preview.text
    for text in ("Некорректный телефон", "Некорректный email", "Выберите значение из списка", "допустимо «Физлицо»"):
        assert text in preview.text
    with factory() as s:  # предпросмотр ничего не пишет
        assert s.scalars(select(Client)).all() == []

    done = confirm(client, "clients", token_of(preview))
    assert done.status_code == 303 and done.headers["location"] == "/clients"
    assert "создано: 2" in client.get("/clients").text
    with factory() as s:
        clients = {c.name: c for c in s.scalars(select(Client))}
        assert set(clients) == {"Анна", "ООО Ромашка"}  # дубль по телефону пропущен
        assert clients["Анна"].custom[source] == "Сайт" and clients["Анна"].owner_id is not None
        assert clients["ООО Ромашка"].type.value == "company"
    # токен одноразовый
    again = confirm(client, "clients", token_of(preview))
    assert again.status_code == 303 and again.headers["location"] == "/import/clients"


def test_clients_import_update_mode_does_not_erase_with_blank_cells(client, make_user, factory):
    login_admin(client, make_user)
    client.post(
        "/api/v1/clients",
        json={"name": "Анна", "email": "anna@example.com", "phone": "+7 916 111-22-33",
              "company_name": "Старая", "address": "Москва"},
    )
    data = csv_file([["Имя", "Email", "Компания", "Адрес"], ["Анна Новая", "ANNA@example.com", "Новая", ""]])

    skipped = upload(client, "clients", "c.csv", data)  # по умолчанию повторы пропускаются
    assert 'name="token"' not in skipped.text and "уже есть" in skipped.text

    preview = upload(client, "clients", "c.csv", data, duplicates="update")
    assert "Ошибки" not in preview.text
    confirm(client, "clients", token_of(preview))
    with factory() as s:
        (anna,) = s.scalars(select(Client)).all()
        assert (anna.name, anna.company_name, anna.address) == ("Анна Новая", "Новая", "Москва")
        assert anna.phone == "+7 916 111-22-33"


def test_clients_import_from_excel_with_native_cell_types(client, make_user, factory):
    login_admin(client, make_user)
    with factory() as s:
        codes = {
            label: custom_field_service.create_field(
                s, entity_type="client", label=label, field_type=ftype, options=opts, code=code
            ).code
            for label, ftype, opts, code in [
                ("Бюджет", "decimal", [], "budget"), ("Дата знакомства", "date", [], "met"),
                ("Согласие", "boolean", [], "agree"), ("Теги", "multiselect", ["a", "b"], "tags"),
                ("Визит", "datetime", [], "visit"),
            ]
        }
    rows = [
        ["ФИО", "Мобильный", "Бюджет", "Дата знакомства", "Согласие", "Теги", "Визит"],
        ["Борис", 79161234567, 1500.5, datetime(2026, 3, 5), "да", "a; b", datetime(2026, 3, 5, 14, 30)],
    ]
    preview = upload(client, "clients", "clients.xlsx", xlsx_file(rows))
    assert "Ошибки" not in preview.text
    confirm(client, "clients", token_of(preview))
    with factory() as s:
        (boris,) = s.scalars(select(Client)).all()
        assert boris.phone == "79161234567"
        custom = boris.custom
        assert custom[codes["Бюджет"]] == Decimal("1500.5")
        assert custom[codes["Дата знакомства"]] == date(2026, 3, 5)
        assert custom[codes["Согласие"]] is True and custom[codes["Теги"]] == ["a", "b"]
        assert to_local(custom[codes["Визит"]]).replace(tzinfo=None) == datetime(2026, 3, 5, 14, 30)


def test_import_file_level_problems(client, make_user):
    login_admin(client, make_user)
    no_name = upload(client, "clients", "a.csv", csv_file([["Телефон"], ["+7 916 111-22-33"]]))
    assert no_name.status_code == 200 and "нет обязательных колонок: «Имя»" in no_name.text
    assert 'name="token"' not in no_name.text

    unknown = upload(client, "clients", "a.csv", csv_file([["Имя", "Цвет"], ["Анна", "красный"]]))
    assert "Колонка «Цвет» не распознана" in unknown.text and token_of(unknown)

    assert upload(client, "clients", "a.pdf", b"%PDF").status_code == 400
    assert upload(client, "clients", "a.xlsx", b"not a zip").status_code == 400
    old = upload(client, "clients", "a.xls", b"x")
    assert old.status_code == 400 and "Старый формат" in old.text
    assert client.post("/import/clients", data={}).status_code == 400  # файл не выбран

    for bad in ("0" * 32, "../../etc/passwd", ""):
        resp = confirm(client, "clients", bad)
        assert resp.status_code == 303 and resp.headers["location"] == "/import/clients"
    assert client.get("/import/unknown").status_code == 404


def test_confirm_works_only_for_the_user_who_uploaded(client, make_user, factory):
    make_user("a@example.com", role="admin", full_name="А")
    make_user("b@example.com", role="admin", full_name="Б")
    web_login(client, "a@example.com")
    token = token_of(upload(client, "clients", "c.csv", csv_file([["Имя"], ["Анна"]])))
    web_login(client, "b@example.com")
    resp = confirm(client, "clients", token)
    assert resp.headers["location"] == "/import/clients"
    with factory() as s:
        assert s.scalars(select(Client)).all() == []
    web_login(client, "a@example.com")
    assert confirm(client, "clients", token).headers["location"] == "/clients"
    with factory() as s:
        assert [c.name for c in s.scalars(select(Client))] == ["Анна"]


# ------------------------------------------------------------------ сделки
def test_deals_import_links_clients_statuses_and_templates(client, make_user, factory):
    login_admin(client, make_user)
    with factory() as s:
        presets.apply_preset(s, "auto")
    client.post("/api/v1/clients", json={"name": "Анна", "phone": "+7 (916) 111-22-33", "email": "anna@example.com"})
    client.post("/api/v1/clients", json={"name": "Борис", "email": "boris@example.com"})
    rows = [
        ["Название", "Телефон клиента", "Email клиента", "Сумма", "Дата", "Статус", "Шаблон", "Марка авто", "VIN"],
        ["Замена масла", "8 916 111-22-33", "", "1 500,50", "05.03.2026", "Приём", "Автосервис", "Toyota", "XW1"],
        ["Диагностика", "", "boris@example.com", "3000", "", "Диагностика", "Автосервис", "", ""],
        ["Без клиента", "", "", "100", "", "", "", "", ""],
        ["Чужой телефон", "000 999 88 77", "", "100", "", "", "", "", ""],
        ["Поле без шаблона", "", "anna@example.com", "10", "", "", "", "Honda", ""],
        ["Левый статус", "", "anna@example.com", "10", "", "Нет такого", "", "", ""],
        ["Стандарт", "", "anna@example.com", "", "", "", "Стандартная форма", "", ""],
    ]
    preview = upload(client, "deals", "deals.csv", csv_file(rows))
    assert preview.status_code == 200 and "Ошибки (4)" in preview.text
    for text in (
        "Укажите телефон или email клиента",
        "не найден — сначала загрузите клиентов",
        "поле есть только у шаблона «Автосервис»",
        "Статус: «Нет такого» не найден в настройках",
    ):
        assert text in preview.text
    with factory() as s:
        assert s.scalars(select(Deal)).all() == []

    confirm(client, "deals", token_of(preview))
    with factory() as s:
        deals = {d.title: d for d in s.scalars(select(Deal))}
        assert set(deals) == {"Замена масла", "Диагностика", "Стандарт"}
        oil = deals["Замена масла"]
        assert oil.client.name == "Анна"  # 8 916… и +7 916… — один телефон
        assert (oil.amount, oil.deal_date) == (Decimal("1500.50"), date(2026, 3, 5))
        assert oil.status.name == "Приём" and oil.template.name == "Автосервис"
        assert {"Toyota", "XW1"} <= set(oil.custom.values())
        assert deals["Диагностика"].client.name == "Борис" and deals["Диагностика"].status.name == "Диагностика"
        assert deals["Стандарт"].template is None and deals["Стандарт"].amount == 0


def test_deals_import_skips_repeats_unless_allowed(client, make_user, factory):
    login_admin(client, make_user)
    client.post("/api/v1/clients", json={"name": "Анна", "email": "anna@example.com"})
    rows = [
        ["Название", "Email клиента", "Сумма", "Дата"],
        ["Замена масла", "anna@example.com", "1500", "05.03.2026"],
        ["ЗАМЕНА МАСЛА", "anna@example.com", "1 500,00", "05.03.2026"],  # тот же заказ в файле дважды
        ["Замена масла", "anna@example.com", "1600", "05.03.2026"],      # другая сумма — другая запись
    ]
    data = csv_file(rows)

    first = upload(client, "deals", "d.csv", data)
    assert "Такая запись уже есть" in first.text
    confirm(client, "deals", token_of(first))
    with factory() as s:
        assert len(s.scalars(select(Deal)).all()) == 2

    # повторная загрузка того же файла ничего не добавляет
    again = upload(client, "deals", "d.csv", data)
    assert 'name="token"' not in again.text
    with factory() as s:
        assert len(s.scalars(select(Deal)).all()) == 2

    # осознанный режим «загрузить всё равно»
    forced = upload(client, "deals", "d.csv", data, duplicates="allow")
    confirm(client, "deals", token_of(forced))
    with factory() as s:
        assert len(s.scalars(select(Deal)).all()) == 5


def test_deals_import_requires_a_client_column(client, make_user):
    login_admin(client, make_user)
    resp = upload(client, "deals", "d.csv", csv_file([["Название", "Сумма"], ["Заказ", "100"]]))
    assert "Нужна колонка, по которой находится клиент" in resp.text and 'name="token"' not in resp.text


def test_employee_import_cannot_assign_and_exports_only_own_deals(client, make_user, factory):
    employee_id = make_user("emp@example.com", role="employee", full_name="Ольга")
    make_user("mgr@example.com", role="manager", full_name="Мила")
    web_login(client, "mgr@example.com")
    cid = client.post("/api/v1/clients", json={"name": "Анна", "email": "anna@example.com"}).json()["id"]
    client.post("/api/v1/deals", json={"title": "Менеджерская", "client_id": cid})

    web_login(client, "emp@example.com")
    rows = [["Название", "Email клиента", "Ответственный"], ["Моя", "anna@example.com", ""], ["Чужая", "anna@example.com", "mgr@example.com"]]
    preview = upload(client, "deals", "d.csv", csv_file(rows))
    assert "Нет права назначать" in preview.text and "Ошибки (1)" in preview.text
    confirm(client, "deals", token_of(preview))
    with factory() as s:
        mine = {d.title: d for d in s.scalars(select(Deal))}
        assert mine["Моя"].responsible_id == employee_id and "Чужая" not in mine

    exported = client.get("/deals/export?format=csv")
    text = exported.content.decode("utf-8-sig")
    assert "Моя" in text and "Менеджерская" not in text

    web_login(client, "mgr@example.com")
    text = client.get("/deals/export?format=csv").content.decode("utf-8-sig")
    assert "Моя" in text and "Менеджерская" in text


# ------------------------------------------------------------------ экспорт и шаблоны
def sheet_rows(content):
    ws = load_workbook(io.BytesIO(content), read_only=True).worksheets[0]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_export_respects_filters_and_includes_custom_fields(client, make_user, factory):
    login_admin(client, make_user)
    with factory() as s:
        custom_field_service.create_field(s, entity_type="deal", label="Гарантия", field_type="text", code="warranty")
        custom_field_service.create_field(s, entity_type="client", label="Источник", field_type="text", code="source")
    anna = client.post("/api/v1/clients", json={"name": "Анна", "custom": {"source": "Сайт"}}).json()["id"]
    client.post("/api/v1/clients", json={"name": "ООО Ромашка", "type": "company"})
    client.post("/api/v1/deals", json={"title": "Первая", "client_id": anna, "amount": "100", "custom": {"warranty": "12 мес"}})
    client.post("/api/v1/deals", json={"title": "Вторая", "client_id": anna, "amount": "200"})

    resp = client.get("/deals/export?format=xlsx&q=Перв")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert re.fullmatch(r'attachment; filename="deals-\d{4}-\d{2}-\d{2}\.xlsx"', resp.headers["content-disposition"])
    rows = sheet_rows(resp.content)
    assert rows[0][:3] == ["Название", "Клиент", "Телефон клиента"] and rows[0][-1] == "Гарантия"
    assert len(rows) == 2 and rows[1][0] == "Первая" and rows[1][1] == "Анна" and rows[1][-1] == "12 мес"
    assert rows[1][rows[0].index("Сумма")] == 100

    csv_all = client.get("/deals/export?format=csv").content.decode("utf-8-sig")
    assert "Первая" in csv_all and "Вторая" in csv_all and csv_all.splitlines()[0].startswith("Название;Клиент;")

    companies = sheet_rows(client.get("/clients/export?type=company").content)  # xlsx по умолчанию
    assert [r[1] for r in companies[1:]] == ["ООО Ромашка"] and companies[0][-1] == "Источник"
    assert companies[1][0] == "Компания"

    page = client.get("/clients")  # в списке есть меню с учётом текущих фильтров
    assert "/clients/export?format=xlsx" in page.text and "/import/clients" in page.text
    assert "type=company&amp;format=csv" in client.get("/clients?type=company").text


def test_templates_have_headers_for_every_field(client, make_user, factory):
    login_admin(client, make_user)
    with factory() as s:
        presets.apply_preset(s, "auto")
        presets.apply_preset(s, "beauty")
    resp = client.get("/import/deals/template?format=xlsx")
    wb = load_workbook(io.BytesIO(resp.content))
    headers = [c.value for c in wb["Данные"][1]]
    for expected in ("Название", "Телефон клиента", "Email клиента", "Статус", "Шаблон", "Марка авто", "VIN", "Мастер", "Услуга"):
        assert expected in headers
    assert wb["Подсказки"].max_row == len(headers) + 1 and wb["Данные"].data_validations.dataValidation
    hints = {row[0].value: row[2].value for row in wb["Подсказки"].iter_rows(min_row=2)}
    assert "Только для шаблона «Автосервис»" in hints["VIN"] and "Анна" in hints["Мастер"]

    csv_text = client.get("/import/clients/template?format=csv").content.decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Тип;Имя;Компания;Email;Телефон;")

    # скачанный шаблон сразу пригоден: пустой файл проверяется без ошибок формата
    assert upload(client, "deals", "t.xlsx", resp.content).status_code == 200


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_export_then_import_keeps_every_field_type(client, make_user, factory, fmt):
    login_admin(client, make_user)
    with factory() as s:
        for label, ftype, opts, code in [
            ("Заметка", "text", [], "note"), ("Этаж", "integer", [], "floor"), ("Бюджет", "decimal", [], "budget"),
            ("День", "date", [], "day"), ("Визит", "datetime", [], "visit"), ("Согласие", "boolean", [], "agree"),
            ("Канал", "select", ["Сайт", "Звонок"], "channel"), ("Теги", "multiselect", ["a", "b", "в"], "tags"),
        ]:
            custom_field_service.create_field(s, entity_type="client", label=label, field_type=ftype, options=opts, code=code)
    created = client.post(
        "/api/v1/clients",
        json={
            "name": "Анна", "phone": "+7 (916) 111-22-33", "email": "anna@example.com", "company_name": "ООО",
            "address": "Москва, ул. Ленина 1", "notes": 'Строка 1\nСтрока 2; с запятой, и "кавычками"',
            "custom": {
                "note": "привет", "floor": 5, "budget": "12.5", "day": "2026-01-02", "visit": "2026-01-02T10:30",
                "agree": True, "channel": "Сайт", "tags": ["a", "в"],
            },
        },
    )
    assert created.status_code == 201, created.text
    before = client.get(f"/api/v1/clients/{created.json()['id']}").json()  # как лежит в БД, а не «сырой» ответ создания
    exported = client.get(f"/clients/export?format={fmt}")
    assert client.delete(f"/api/v1/clients/{before['id']}").status_code == 204

    preview = upload(client, "clients", f"clients.{fmt}", exported.content)
    assert "Ошибки" not in preview.text
    confirm(client, "clients", token_of(preview))
    (after,) = client.get("/api/v1/clients").json()["items"]
    for key in ("name", "phone", "email", "company_name", "address", "notes", "type", "custom"):
        assert after[key] == before[key], key


# ------------------------------------------------------------------ права и защита
def test_permissions(client, make_user, factory):
    make_user("admin@example.com", role="admin")
    with factory() as s:
        role_service.create_role(
            s, code="viewer", name="Просмотр", description=None, permissions=["clients:read", "deals:read"]
        )
    make_user("viewer@example.com", role="viewer")
    assert client.get("/import/clients", follow_redirects=False).status_code in (302, 303)  # без входа
    assert client.get("/clients/export").status_code in (302, 303)

    web_login(client, "viewer@example.com")
    assert client.get("/import/clients").status_code == 403
    assert client.get("/import/deals/template").status_code == 403
    assert upload(client, "clients", "a.csv", csv_file([["Имя"], ["Анна"]])).status_code == 403
    assert confirm(client, "clients", "0" * 32).status_code == 403
    assert client.get("/clients/export").status_code == 200  # просмотр даёт выгрузку
    assert client.get("/deals/export?format=csv").status_code == 200
    page = client.get("/clients").text
    assert "/clients/export" in page and "/import/clients" not in page  # пункта «Загрузить» нет


def test_upload_needs_csrf_token_when_protection_is_on(client, make_user, monkeypatch):
    login_admin(client, make_user)  # вход выполняется до включения защиты
    monkeypatch.setattr(get_settings(), "CSRF_ENABLED", True)
    data = csv_file([["Имя"], ["Анна"]])
    assert upload(client, "clients", "a.csv", data).status_code == 403
    form = client.get("/import/clients").text
    token = re.search(r'name="csrf_token" value="([^"]+)"', form).group(1)
    ok = client.post("/import/clients", files={"file": ("a.csv", data)}, data={"csrf_token": token})
    assert ok.status_code == 200 and 'name="token"' in ok.text


def test_upload_page_renders_with_template_links(client, make_user):
    login_admin(client, make_user)
    page = client.get("/import/deals").text
    assert "Загрузка из файла" in page and "/import/deals/template?format=csv" in page
