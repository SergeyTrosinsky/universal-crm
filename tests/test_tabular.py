"""Чтение и запись CSV / XLSX и сопоставление колонок (без базы данных)."""
import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.models import CustomField, FieldType
from app.services import import_export_service as ie
from app.services import tabular
from app.services.tabular import FileError, XlsxSpec, read_table, write_csv, write_xlsx


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def field(fid, label, code, ftype=FieldType.TEXT, **extra):
    return CustomField(id=fid, label=label, code=code, field_type=ftype, options=extra.pop("options", []), **extra)


def test_csv_delimiters_and_encodings():
    for delimiter in (";", ",", "\t"):
        raw = f"Имя{delimiter}Телефон\r\nОльга{delimiter}12345\r\n".encode("utf-8-sig")
        table = read_table(raw, "a.csv", max_rows=10)
        assert table.headers == ["Имя", "Телефон"] and table.rows == [(2, ["Ольга", "12345"])]
    table = read_table("Имя;Город\nОльга;Тверь\n".encode("cp1251"), "a.csv", max_rows=10)
    assert table.rows[0][1] == ["Ольга", "Тверь"]


def test_blank_rows_and_cells_are_skipped_and_numbers_normalised():
    table = read_table(xlsx([["Имя", "Этаж", None], [None, None, None], ["Анна", 5.0, None], ["Борис", 2.5, None]]), "a.xlsx", max_rows=10)
    assert table.headers == ["Имя", "Этаж"]
    assert [row for _, row in table.rows] == [["Анна", 5, None], ["Борис", 2.5, None]]
    assert [n for n, _ in table.rows] == [3, 4]


@pytest.mark.parametrize(
    "data, name, text",
    [
        (b"%PDF", "a.pdf", "Поддерживаются файлы"),
        (b"x", "a.xls", "Старый формат"),
        (b"", "a.csv", "Файл пустой"),
        (b"not a zip", "a.xlsx", "Не удалось прочитать файл Excel"),
        (b"\x00\x01\x02\x03", "a.csv", "не текстовый"),
        (b"a;b\n1;2\n1;2\n1;2\n", "a.csv", "Слишком много строк"),
    ],
)
def test_unreadable_files_give_clear_errors(data, name, text):
    with pytest.raises(FileError) as e:
        read_table(data, name, max_rows=2)
    assert text in str(e.value)


def test_csv_roundtrip_and_formula_protection():
    raw = write_csv(
        ["Имя", "Телефон", "Сумма", "Дата"],
        [["=HYPERLINK(1)", "+7 (916) 123-45-67", Decimal("1500.50"), date(2026, 1, 5)], ["@cmd", "-abc", Decimal("3"), None]],
    )
    assert raw.startswith(b"\xef\xbb\xbf") and b";" in raw
    table = read_table(raw, "a.csv", max_rows=10)
    assert table.rows[0][1] == ["=HYPERLINK(1)", "+7 (916) 123-45-67", "1500,50", "05.01.2026"]
    assert table.rows[1][1] == ["@cmd", "-abc", "3", None]
    assert raw.decode("utf-8-sig").splitlines()[1].startswith("'=HYPERLINK(1);")


def test_xlsx_cells_are_typed_and_formulas_are_text():
    raw = write_xlsx(
        XlsxSpec(
            headers=["Имя", "Сумма", "Дата", "Визит"],
            rows=[["=1+1", Decimal("10.5"), date(2026, 1, 2), datetime(2026, 1, 2, 10, 30)]],
            hints=[("Имя", "да", "текст")],
            lists=[tabular.ListValidation(1, ["Анна", "Борис"])],
            money_columns={2},
        )
    )
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Данные"]
    assert ws["A2"].data_type == "s" and ws["A2"].value == "'=1+1"
    assert ws["B2"].value == 10.5 and ws["B2"].number_format == "#,##0.00"
    assert ws["C2"].number_format == "DD.MM.YYYY" and ws["D2"].number_format == "DD.MM.YYYY HH:MM"
    assert wb["Подсказки"]["A2"].value == "Имя"
    assert wb["Списки"].sheet_state == "hidden"
    assert ws.data_validations.dataValidation and ws.freeze_panes == "A2"
    assert read_table(raw, "a.xlsx", max_rows=5).rows[0][1][0] == "=1+1"


def test_columns_are_found_by_aliases_labels_and_codes():
    fields = [field(1, "VIN", "vin"), field(2, "Марка авто", "car_brand")]
    cmap = ie.map_columns(["ФИО", "Мобильный", "vin", "МАРКА  авто", "Что-то"], ie.CLIENT_COLUMNS, fields)
    assert cmap.base == {"name": 0, "phone": 1}
    assert cmap.custom == {1: 2, 2: 3}
    assert cmap.warnings == ["Колонка «Что-то» не распознана и будет пропущена"]


def test_colliding_and_duplicate_labels_get_a_code_suffix():
    fields = [field(1, "Телефон", "second_phone", FieldType.PHONE), field(2, "Мастер", "m1"), field(3, "Мастер", "m2")]
    names = ie.field_headers(fields, ie._reserved(ie.CLIENT_COLUMNS))
    assert names == {1: "Телефон (second_phone)", 2: "Мастер (m1)", 3: "Мастер (m2)"}
    cmap = ie.map_columns(["Телефон", "Телефон (second_phone)", "Мастер (m2)", "Мастер"], ie.CLIENT_COLUMNS, fields)
    assert cmap.base == {"phone": 0}
    assert cmap.custom == {1: 1, 3: 2}
    assert cmap.warnings == ["Колонка «Мастер» не распознана и будет пропущена"]


def test_repeated_header_uses_first_column():
    cmap = ie.map_columns(["Имя", "Email", "email"], ie.CLIENT_COLUMNS, [])
    assert cmap.base == {"name": 0, "email": 1}
    assert "повторяется" in cmap.warnings[0]


def test_phone_and_email_keys_ignore_formatting():
    assert ie.phone_key("+7 (916) 111-22-33") == ie.phone_key("8 916 111 22 33") == ie.phone_key(89161112233) == "9161112233"
    assert ie.phone_key("12") is None and ie.phone_key("") is None
    assert ie.email_key(" Anna@Example.COM ") == "anna@example.com" and ie.email_key("x@@bad") is None
